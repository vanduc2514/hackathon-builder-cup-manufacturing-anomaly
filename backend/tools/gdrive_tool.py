"""ADK Tool: Cloud-native Google Drive API v3 connector for external knowledge base access."""

from __future__ import annotations

import asyncio
import io
import logging
import os
from pathlib import Path
from typing import Any, Dict, List, Optional

from backend.core.config import settings
from backend.core.models import GoogleDriveDocument

logger = logging.getLogger(__name__)

# Standard MIME types
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_GOOGLE_DOC = "application/vnd.google-apps.document"
MIME_GOOGLE_SHEET = "application/vnd.google-apps.spreadsheet"

# Fallback mock documents for offline / zero-credential local dev parity
MOCK_GDRIVE_DOCUMENTS: List[GoogleDriveDocument] = [
    GoogleDriveDocument(
        file_id="mock-doc-cnc-01",
        web_view_link="https://drive.google.com/file/d/mock-doc-cnc-01/view",
        path="manuals/cnc_operating_manual.docx",
        name="cnc_operating_manual.docx",
        size_bytes=15420,
        is_dir=False,
        mime_type=MIME_DOCX,
        mod_time="2026-03-12T10:00:00Z",
    ),
    GoogleDriveDocument(
        file_id="mock-doc-cnc-maint-01",
        web_view_link="https://drive.google.com/file/d/mock-doc-cnc-maint-01/view",
        path="maint_logs/cnc_tool_t07_maintenance.xlsx",
        name="cnc_tool_t07_maintenance.xlsx",
        size_bytes=8400,
        is_dir=False,
        mime_type=MIME_XLSX,
        mod_time="2026-03-12T11:00:00Z",
    ),
    GoogleDriveDocument(
        file_id="mock-doc-im-01",
        web_view_link="https://drive.google.com/file/d/mock-doc-im-01/view",
        path="manuals/injection_molding_guide.docx",
        name="injection_molding_guide.docx",
        size_bytes=12800,
        is_dir=False,
        mime_type=MIME_DOCX,
        mod_time="2026-03-12T12:00:00Z",
    ),
]

MOCK_DOC_TEXTS: Dict[str, str] = {
    "mock-doc-cnc-01": (
        "# CNC VMC-850 Operating Manual\n\n"
        "## Section 4.2 - Spindle Tooling & Insert Replacement\n"
        "When machining stainless and alloy steels, tool station T07 utilizes insert WNMG-080408. "
        "Excessive tool flank wear (VB > 0.3mm) causes vertical scratch defects on workpiece outer surfaces. "
        "Upon insert replacement, torque clamping screw to 4.2 N·m. "
        "Measure spindle and arbor radial runout with dial indicator. Runout must not exceed 0.005 mm (tolerance limit).\n"
    ),
    "mock-doc-cnc-maint-01": (
        "### Sheet: Tooling Maintenance Checklist\n"
        "Station: T07 | Insert_Code: WNMG-080408 | Torque_Spec: 4.2 N·m | Max_Runout: 0.005 mm | Status: Active\n"
        "Station: T08 | Insert_Code: CNMG-120408 | Torque_Spec: 5.0 N·m | Max_Runout: 0.005 mm | Status: Active\n"
    ),
    "mock-doc-im-01": (
        "# Injection Molding IM-05 Operating Guide\n\n"
        "## Section 3.1 - Resin Dehumidifier & Dryer System\n"
        "Short fill and sink mark defects occur when desiccant canister is saturated or dewpoint exceeds -30.0 °C. "
        "Replace auxiliary dryer filter with FLT-ABS-04. Verify dewpoint reaches -35.0 °C or lower.\n"
    ),
}


class GoogleDriveConnector:
    """Cloud-native Google Drive API v3 connector.

    Replaces rclone CLI wrappers with official google-api-python-client.
    Authenticates automatically via Application Default Credentials (ADC) in Cloud Run
    and provides offline mock fallbacks for local dev parity.
    """

    def __init__(
        self,
        folder_id: Optional[str] = None,
        service_account_path: Optional[str] = None,
        cache_dir: Optional[str | Path] = None,
        timeout_seconds: float = 30.0,
    ) -> None:
        self.folder_id = folder_id or settings.gdrive_folder_id
        self.service_account_path = service_account_path or settings.gdrive_service_account_path
        self.cache_dir = Path(cache_dir or settings.gdrive_cache_dir)
        self.timeout = timeout_seconds
        self._service: Optional[Any] = None
        self._initialized = False

    def _get_service(self) -> Optional[Any]:
        """Lazy-initialize Google Drive v3 resource using Application Default Credentials."""
        if self._initialized:
            return self._service

        self._initialized = True
        try:
            from google.auth import default
            from google.oauth2 import service_account
            from googleapiclient.discovery import build

            scopes = ["https://www.googleapis.com/auth/drive.readonly"]

            if self.service_account_path and os.path.isfile(self.service_account_path):
                creds = service_account.Credentials.from_service_account_file(
                    self.service_account_path, scopes=scopes
                )
            else:
                creds, _ = default(scopes=scopes)

            self._service = build("drive", "v3", credentials=creds, cache_discovery=False)
            logger.info("Successfully initialized Google Drive v3 API client")
            return self._service
        except Exception as exc:
            logger.debug(
                "Google Drive API v3 credentials unavailable (%s); falling back to mock mode",
                exc,
            )
            self._service = None
            return None

    def is_available(self) -> bool:
        """Check if live Google Drive v3 client is connected and accessible."""
        return self._get_service() is not None

    def list_documents(
        self, folder_id: Optional[str] = None, recursive: bool = False
    ) -> List[GoogleDriveDocument]:
        """List files in the target Google Drive folder."""
        service = self._get_service()
        target_folder = folder_id or self.folder_id

        if not service or not target_folder:
            logger.debug("Listing Google Drive documents from local mock fallback")
            return list(MOCK_GDRIVE_DOCUMENTS)

        try:
            query = f"'{target_folder}' in parents and trashed = false"
            fields = "files(id, name, mimeType, size, modifiedTime, webViewLink)"
            results = (
                service.files()
                .list(
                    q=query,
                    fields=fields,
                    supportsAllDrives=True,
                    includeItemsFromAllDrives=True,
                    pageSize=100,
                )
                .execute()
            )

            files = results.get("files", [])
            documents: List[GoogleDriveDocument] = []
            for f in files:
                documents.append(
                    GoogleDriveDocument(
                        file_id=f.get("id"),
                        web_view_link=f.get("webViewLink"),
                        path=f.get("name", ""),
                        name=f.get("name", ""),
                        size_bytes=int(f.get("size", 0)),
                        is_dir=f.get("mimeType") == "application/vnd.google-apps.folder",
                        mime_type=f.get("mimeType"),
                        mod_time=f.get("modifiedTime"),
                    )
                )
            return documents
        except Exception as exc:
            logger.warning("Failed to list Google Drive files via v3 API: %s", exc)
            return list(MOCK_GDRIVE_DOCUMENTS)

    def stream_document_bytes(self, file_id: str) -> bytes:
        """Download raw binary content of a file from Google Drive into memory."""
        service = self._get_service()
        if not service:
            text = MOCK_DOC_TEXTS.get(file_id, "Mock document content for offline parity.")
            return text.encode("utf-8")

        try:
            request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
            return request.execute()
        except Exception as exc:
            logger.warning("Error streaming document bytes for file_id '%s' (%s); falling back to mock content", file_id, exc)
            text = MOCK_DOC_TEXTS.get(file_id, "Mock document content for offline parity.")
            return text.encode("utf-8")

    def export_google_doc_text(self, file_id: str, mime_type: str) -> str:
        """Export native Google Docs or Sheets to text / CSV format."""
        service = self._get_service()
        if not service:
            return MOCK_DOC_TEXTS.get(file_id, "")

        export_mime = "text/plain"
        if mime_type == MIME_GOOGLE_SHEET:
            export_mime = "text/csv"

        try:
            request = service.files().export_media(fileId=file_id, mimeType=export_mime)
            content_bytes = request.execute()
            return content_bytes.decode("utf-8", errors="replace")
        except Exception as exc:
            logger.error("Error exporting Google Doc '%s': %s", file_id, exc)
            raise RuntimeError(f"Failed to export Google Doc '{file_id}': {exc}") from exc

    def extract_document_text(
        self,
        file_bytes: bytes,
        mime_type: Optional[str] = None,
        filename: Optional[str] = None,
    ) -> str:
        """Format-agnostic in-memory document parser converting docx/xlsx/text into Markdown."""
        name = (filename or "").lower()
        mime = (mime_type or "").lower()

        # Handle Word (.docx)
        if name.endswith(".docx") or mime == MIME_DOCX:
            return self._parse_docx_in_memory(file_bytes)

        # Handle Excel (.xlsx)
        if name.endswith(".xlsx") or mime == MIME_XLSX:
            return self._parse_xlsx_in_memory(file_bytes)

        # Handle plain text / markdown / csv
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="replace")

    @staticmethod
    def _parse_docx_in_memory(file_bytes: bytes) -> str:
        """Extract paragraphs and tables from .docx bytes into clean GitHub-flavored Markdown."""
        try:
            from docx import Document

            doc = Document(io.BytesIO(file_bytes))
            sections: List[str] = []

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                # Preserve heading hierarchy if styled
                style_name = getattr(p.style, "name", "").lower()
                if "heading 1" in style_name or "title" in style_name:
                    sections.append(f"# {text}")
                elif "heading 2" in style_name:
                    sections.append(f"## {text}")
                elif "heading 3" in style_name:
                    sections.append(f"### {text}")
                else:
                    sections.append(text)

            # Preserve tables as markdown tables
            for table in doc.tables:
                table_lines: List[str] = []
                for row_idx, row in enumerate(table.rows):
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if not any(cells):
                        continue
                    if row_idx == 0:
                        table_lines.append("| " + " | ".join(cells) + " |")
                        table_lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
                    else:
                        table_lines.append("| " + " | ".join(cells) + " |")
                if table_lines:
                    sections.append("\n".join(table_lines))

            return "\n\n".join(sections)
        except Exception as exc:
            logger.warning("docx parsing failed; falling back to raw decode: %s", exc)
            return file_bytes.decode("utf-8", errors="replace")

    @staticmethod
    def _parse_xlsx_in_memory(file_bytes: bytes) -> str:
        """Extract Excel sheets into header-anchored key-value text records."""
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            sheet_sections: List[str] = []

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue

                # Locate first non-empty row as column headers
                header_idx = -1
                headers: List[str] = []
                for idx, r in enumerate(rows):
                    non_empty = [
                        str(cell).strip() for cell in r if cell is not None and str(cell).strip()
                    ]
                    if non_empty:
                        header_idx = idx
                        headers = [
                            str(cell).strip() if cell is not None else f"Col_{i}"
                            for i, cell in enumerate(r)
                        ]
                        break

                if header_idx == -1:
                    continue

                records: List[str] = [f"### Sheet: {sheet_name}"]
                for r in rows[header_idx + 1 :]:
                    if not any(cell is not None for cell in r):
                        continue
                    record_parts = []
                    for h, val in zip(headers, r):
                        if val is not None and str(val).strip():
                            record_parts.append(f"{h}: {val}")
                    if record_parts:
                        records.append(" | ".join(record_parts))

                if len(records) > 1:
                    sheet_sections.append("\n".join(records))

            return "\n\n".join(sheet_sections)
        except Exception as exc:
            logger.warning("xlsx parsing failed; falling back to raw decode: %s", exc)
            return file_bytes.decode("utf-8", errors="replace")

    def search_exact_documents(
        self, query_terms: List[str], folder_id: Optional[str] = None
    ) -> List[GoogleDriveDocument]:
        """Search Google Drive for files containing exact terms via fullText contains."""
        service = self._get_service()
        target_folder = folder_id or self.folder_id

        if not service or not target_folder:
            # Local mock search
            matches: List[GoogleDriveDocument] = []
            for doc in MOCK_GDRIVE_DOCUMENTS:
                content = MOCK_DOC_TEXTS.get(doc.file_id or "", "")
                if any(
                    term.lower() in content.lower() or term.lower() in doc.name.lower()
                    for term in query_terms
                ):
                    matches.append(doc)
            return matches

        try:
            # Build fullText contains query
            sub_clauses = [
                f"fullText contains '{t}' or name contains '{t}'" for t in query_terms if t
            ]
            if not sub_clauses:
                return self.list_documents(folder_id=target_folder)

            q_clause = " or ".join(sub_clauses)
            query = f"'{target_folder}' in parents and trashed = false and ({q_clause})"
            fields = "files(id, name, mimeType, size, modifiedTime, webViewLink)"

            results = (
                service.files()
                .list(
                    q=query,
                    fields=fields,
                    supportsAllDrives=True,
                    includeItemsFromAllDrives=True,
                    pageSize=20,
                )
                .execute()
            )

            files = results.get("files", [])
            matched: List[GoogleDriveDocument] = []
            for f in files:
                matched.append(
                    GoogleDriveDocument(
                        file_id=f.get("id"),
                        web_view_link=f.get("webViewLink"),
                        path=f.get("name", ""),
                        name=f.get("name", ""),
                        size_bytes=int(f.get("size", 0)),
                        is_dir=f.get("mimeType") == "application/vnd.google-apps.folder",
                        mime_type=f.get("mimeType"),
                        mod_time=f.get("modifiedTime"),
                    )
                )
            return matched
        except Exception as exc:
            logger.warning("Google Drive exact fullText search failed (%s); falling back to mock search", exc)
            matches: List[GoogleDriveDocument] = []
            for doc in MOCK_GDRIVE_DOCUMENTS:
                content = MOCK_DOC_TEXTS.get(doc.file_id or "", "")
                if any(
                    term.lower() in content.lower() or term.lower() in doc.name.lower()
                    for term in query_terms
                ):
                    matches.append(doc)
            return matches

    def fetch_to_cache(self, file_id: str, filename: str, force_refresh: bool = False) -> Path:
        """Optionally fetch remote document to local cache directory (.cache/gdrive)."""
        dest_path = self.cache_dir / filename
        if dest_path.is_file() and not force_refresh:
            return dest_path

        dest_path.parent.mkdir(parents=True, exist_ok=True)
        content = self.stream_document_bytes(file_id)
        dest_path.write_bytes(content)
        return dest_path

    # -------------------------------------------------------------------------
    # Asynchronous methods
    # -------------------------------------------------------------------------

    async def list_documents_async(
        self, folder_id: Optional[str] = None, recursive: bool = False
    ) -> List[GoogleDriveDocument]:
        """Asynchronously list remote documents."""
        return await asyncio.to_thread(self.list_documents, folder_id, recursive)

    async def stream_document_bytes_async(self, file_id: str) -> bytes:
        """Asynchronously stream raw document bytes."""
        return await asyncio.to_thread(self.stream_document_bytes, file_id)

    async def extract_document_text_async(
        self,
        file_bytes: bytes,
        mime_type: Optional[str] = None,
        filename: Optional[str] = None,
    ) -> str:
        """Asynchronously extract document text."""
        return await asyncio.to_thread(self.extract_document_text, file_bytes, mime_type, filename)

    async def search_exact_documents_async(
        self, query_terms: List[str], folder_id: Optional[str] = None
    ) -> List[GoogleDriveDocument]:
        """Asynchronously search exact documents."""
        return await asyncio.to_thread(self.search_exact_documents, query_terms, folder_id)

    async def fetch_to_cache_async(
        self, file_id: str, filename: str, force_refresh: bool = False
    ) -> Path:
        """Asynchronously fetch document to cache."""
        return await asyncio.to_thread(self.fetch_to_cache, file_id, filename, force_refresh)


# Module-level singleton
gdrive_connector = GoogleDriveConnector()


def list_gdrive_documents(
    folder_id: Optional[str] = None, recursive: bool = False
) -> List[GoogleDriveDocument]:
    """Helper to list Google Drive documents using default connector."""
    return gdrive_connector.list_documents(folder_id=folder_id, recursive=recursive)


def stream_gdrive_document(file_id: str) -> bytes:
    """Helper to stream Google Drive document bytes."""
    return gdrive_connector.stream_document_bytes(file_id)


def search_gdrive_exact_documents(
    query_terms: List[str], folder_id: Optional[str] = None
) -> List[GoogleDriveDocument]:
    """Helper to search Google Drive documents with exact terms."""
    return gdrive_connector.search_exact_documents(query_terms=query_terms, folder_id=folder_id)
