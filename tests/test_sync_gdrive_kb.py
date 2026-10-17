"""Unit tests for ``scripts/sync_gdrive_kb.py`` (TKT-069 / Issue #267).

Each test exercises one slice of the S1..S5 pipeline with hermetic fixtures written
into a ``tmp_path`` directory so we never touch the real competition cache. The
test file never imports GCP / Vertex SDKs, so it runs on the regular local
pytest configuration (no ``emulator`` marker needed).
"""

from __future__ import annotations

import json
import os
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

# Make ``scripts/`` importable so we can ``import sync_gdrive_kb`` without it living
# in the python package tree.
_REPO_ROOT = Path(__file__).resolve().parent.parent
_SCRIPTS_DIR = _REPO_ROOT / "scripts"
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

from scripts.sync_gdrive_kb import (  # noqa: E402  (path-based import after sys.path tweak)
    ALARM_SYMPTOM_MAP,
    EightDRecord,
    MaintenanceRecord,
    Redactor,
    _classify_image,
    _is_image_file,
    build_alarm_token_set,
    ingest_8d_xlsx,
    ingest_competition_images,
    ingest_maintenance_xlsx,
    ingest_oem_pdf,
    parse_8d_workbook,
    parse_maintenance_workbook,
    render_8d_markdown,
    render_maintenance_markdown,
    update_index_json,
)


# =============================================================================
# Fixtures
# =============================================================================


@pytest.fixture()
def minimal_8d_xlsx(tmp_path: Path) -> Path:
    """Create a minimal 8D XLSX with the columns the parser expects."""

    path = tmp_path / "8d.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "8D_Cases"
    headers = [
        "record_id",
        "case_type",
        "review_date",
        "machine_id",
        "machine_model",
        "incident_timestamp",
        "alarm_code",
        "fault_category",
        "operating_status",
        "problem_description",
        "D0_emergency_response",
        "D1_team_defined",
        "D2_problem_description_5W2H",
        "D3_containment_action",
        "D4_occurrence_root_cause",
        "D4_escape_root_cause",
        "D4_evidence",
        "D5_permanent_corrective_action",
        "D5_action_owner",
        "D5_due_date",
        "D6_implementation_evidence",
        "D6_effectiveness_metric",
        "D6_effectiveness_result",
        "D6_observation_window_days",
        "D7_prevent_recurrence",
        "D8_closure_and_lessons",
        "report_status",
        "data_quality_score_0_100",
        "usable_for_ai_training",
        "synthetic_data_notice",
    ]
    ws.append(headers)
    ws.append(
        [
            "SYN-8D-001",
            "SYNTHETIC_TRAINING_CASE",
            "2026-09-04",
            "CNC-SIM-02",
            "Mitsubishi MDS-E/EH",
            "2026-09-04T09:15:00+07:00",
            53,
            "Drive Alarm 53",
            "FAULT",
            "Vibration rose on spindle; bearing suspected.",
            "stop machine",
            "team formed",
            "5W2H analysis",
            "Quarantine related parts.",
            "Bearing wear.",
            "No trend review.",
            "Trend log.",
            "Replace bearing.",
            "Maintenance engineer",
            "2026-09-09",
            "Bearing replaced.",
            "Vibration < threshold.",
            "Stable for 7 days.",
            7,
            "Add trend review.",
            "Lesson shared.",
            "CLOSED_SYNTHETIC",
            75,
            "NO",
            "Synthetic",
        ]
    )
    ws.append(
        [
            "SYN-8D-002",
            "SYNTHETIC_TRAINING_CASE",
            "2026-09-07",
            "CNC-SIM-03",
            "Mitsubishi MDS-E/EH",
            "2026-09-07T09:15:00+07:00",
            45,
            "Fan Stop",
            "FAULT",
            "Fan stopped.",
            "stop",
            "team",
            "5w2h",
            "Hold affected parts.",
            "Dust in fan.",
            "Inspection missed.",
            "Inspection photos.",
            "Replace fan.",
            "Tech",
            "2026-09-12",
            "Fan replaced.",
            "Airflow ok.",
            "Stable.",
            7,
            "Add PM check.",
            "Lesson.",
            "CLOSED",
            75,
            "NO",
            "Synthetic",
        ]
    )
    wb.save(path)
    return path


@pytest.fixture()
def minimal_maintenance_xlsx(tmp_path: Path) -> Path:
    path = tmp_path / "maint.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Maintenance_Log"
    headers = [
        "maintenance_record_id",
        "record_type",
        "timestamp",
        "machine_id",
        "machine_model",
        "axis_or_unit",
        "maintenance_type",
        "trigger_or_alarm_code",
        "operating_status_before",
        "reported_symptom",
        "inspection_findings",
        "work_performed",
        "parts_replaced",
        "parameter_or_setting_changed",
        "technician_role",
        "work_order_id",
        "downtime_minutes",
        "test_run_result",
        "operating_status_after",
        "follow_up_required",
        "follow_up_due_date",
        "evidence_reference",
        "linked_8d_record_id",
        "data_source_status",
        "notes",
    ]
    ws.append(headers)
    ws.append(
        [
            "ML-TEST-0001",
            "CORRECTIVE",
            "2026-09-04 10:00:00",
            "CNC-SIM-02",
            "Mitsubishi MDS-E/EH",
            "Spindle drive",
            "Alarm 53",
            53.0,
            "FAULT",
            "Spindle vibration alarm.",
            "Bearing suspect.",
            "Replaced bearing in synthetic workflow.",
            "Bearing (simulated)",
            "None",
            "Maintenance technician",
            "WO-TEST-001",
            60.0,
            "Vibration reduced.",
            "RUNNING",
            "Yes",
            "2026-09-08",
            "EVIDENCE-001",
            "SYN-8D-001",
            "SYNTHETIC",
            "Drive-oem alarm.",
        ]
    )
    wb.save(path)
    return path


@pytest.fixture()
def minimal_index_json(tmp_path: Path) -> Path:
    path = tmp_path / "index.json"
    path.write_text(
        json.dumps(
            {
                "version": "1.0.0",
                "documents": [
                    {
                        "document_id": "manuals/cnc_operating_manual.md",
                        "title": "Existing curated CNC manual",
                        "category": "manual",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    return path


@pytest.fixture()
def image_gdrive_layout(tmp_path: Path) -> Path:
    """Build a tiny ``Original Image/`` + ``image error/`` directory layout."""

    gdrive = tmp_path / "gdrive"
    original_dir = gdrive / "Original Image"
    error_dir = gdrive / "image error"
    original_dir.mkdir(parents=True)
    error_dir.mkdir(parents=True)

    # Two originals (jpg with valid header).
    for i in (1, 2):
        (original_dir / f"orig_{i}.jpg").write_bytes(
            b"\xff\xd8\xff\xe0\x00\x10JFIF" + b"\x00" * 8
        )

    # One error_part file (PNG signature, no extension).
    (error_dir / "Error_part_3").write_bytes(
        b"\x89PNG\r\n\x1a\n" + b"\x00" * 16
    )
    # One nested-subdir layout with a single jpg.
    nested = error_dir / "Error_part_9"
    nested.mkdir()
    (nested / "frame.jpg").write_bytes(b"\xff\xd8\xff\xe0\x00\x10JFIF" + b"\x00" * 8)

    return gdrive


# =============================================================================
# Redaction
# =============================================================================


def test_redactor_substitutes_brand_tokens_case_insensitive():
    redactor = Redactor()
    out = redactor.redact("Mitsubishi MDS-E/EH drive tripped at 09:00")
    assert "drive_oem_a" in out
    assert "drive_series_a" in out
    assert "Mitsubishi" not in out
    assert "MDS-E/EH" not in out


def test_redactor_is_idempotent():
    redactor = Redactor()
    once = redactor.redact("Mitsubishi MDS-E/EH")
    twice = redactor.redact(once)
    assert once == twice


def test_redactor_reports_substitution_counts():
    redactor = Redactor()
    redactor.redact("Mitsubishi once")
    redactor.redact("Mitsubishi twice and MDS-E/EH thrice")
    report = redactor.report.to_dict()
    assert report["Mitsubishi"] == 2
    assert report["MDS-E/EH"] == 1


def test_build_alarm_token_set_includes_natural_variants():
    tokens = build_alarm_token_set(["45", "53"])
    assert "alarm 45" in tokens
    assert "alarm_45" in tokens
    assert "45" in tokens
    assert "alarm 53" in tokens


def test_alarm_symptom_map_matches_plan():
    """Tech Lead D4: 45/46/50/51 -> tool_wear; 52/53 -> spindle_vibration."""

    assert ALARM_SYMPTOM_MAP["45"] == "tool_wear"
    assert ALARM_SYMPTOM_MAP["46"] == "tool_wear"
    assert ALARM_SYMPTOM_MAP["50"] == "tool_wear"
    assert ALARM_SYMPTOM_MAP["51"] == "tool_wear"
    assert ALARM_SYMPTOM_MAP["52"] == "spindle_vibration"
    assert ALARM_SYMPTOM_MAP["53"] == "spindle_vibration"


# =============================================================================
# S1 — 8D XLSX parser
# =============================================================================


def test_parse_8d_workbook_returns_two_records(minimal_8d_xlsx: Path):
    records = parse_8d_workbook(minimal_8d_xlsx)
    assert len(records) == 2
    assert records[0].record_id == "SYN-8D-001"
    assert records[0].alarm_code == 53
    assert records[0].fault_category == "Drive Alarm 53"


def test_parse_8d_workbook_missing_sheet_raises(tmp_path: Path):
    path = tmp_path / "wrong.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Other"
    wb.save(path)
    with pytest.raises(ValueError, match="8D_Cases"):
        parse_8d_workbook(path)


def test_render_8d_markdown_redacts_brand_and_lists_alarm_tag():
    record = EightDRecord(
        record_id="SYN-8D-001",
        machine_id="CNC-SIM-02",
        machine_model="Mitsubishi MDS-E/EH",
        incident_timestamp="2026-09-04T09:15:00+07:00",
        alarm_code=53,
        fault_category="Drive Alarm 53",
        operating_status="FAULT",
        problem_description="Mitsubishi alarm tripped.",
        containment_action="Quarantine.",
        root_cause_occurrence="Bearing wear.",
        root_cause_escape="No trend review.",
        evidence="Trend log.",
        permanent_corrective_action="Replace bearing.",
        action_owner="Maintenance engineer",
        due_date="2026-09-09",
        implementation_evidence="Bearing replaced.",
        effectiveness_metric="Vibration < threshold.",
        effectiveness_result="Stable.",
        observation_window_days=7,
        prevent_recurrence="Add trend review.",
        closure_lessons="Lesson shared.",
    )
    redactor = Redactor()
    md = render_8d_markdown(record, redactor=redactor)
    assert "Mitsubishi" not in md
    assert "drive_oem_a" in md
    assert "alarm_53" in md
    assert "D2: Problem Description" in md


def test_ingest_8d_xlsx_writes_markdown_per_record(minimal_8d_xlsx: Path, tmp_path: Path):
    out_dir = tmp_path / "8d_md"
    written = ingest_8d_xlsx(minimal_8d_xlsx, out_dir, redactor=Redactor())
    assert len(written) == 2
    assert (out_dir / "8D_SYN-8D-001.md").exists()
    assert (out_dir / "8D_SYN-8D-002.md").exists()


# =============================================================================
# S2 — Maintenance log parser
# =============================================================================


def test_parse_maintenance_workbook_returns_records(minimal_maintenance_xlsx: Path):
    records = parse_maintenance_workbook(minimal_maintenance_xlsx)
    assert len(records) == 1
    assert records[0].record_id == "ML-TEST-0001"
    assert records[0].linked_8d_record_id == "SYN-8D-001"


def test_render_maintenance_markdown_redacts_brand():
    record = MaintenanceRecord(
        record_id="ML-TEST-0001",
        record_type="CORRECTIVE",
        timestamp="2026-09-04 10:00:00",
        machine_id="CNC-SIM-02",
        machine_model="Mitsubishi MDS-E/EH",
        axis_or_unit="Spindle drive",
        maintenance_type="Alarm 53",
        trigger_alarm_code=53.0,
        operating_status_before="FAULT",
        reported_symptom="Spindle vibration alarm.",
        inspection_findings="Bearing suspect.",
        work_performed="Replaced bearing.",
        parts_replaced="Bearing (simulated)",
        parameter_changed="None",
        technician_role="Maintenance technician",
        work_order_id="WO-TEST-001",
        downtime_minutes=60.0,
        test_run_result="Vibration reduced.",
        operating_status_after="RUNNING",
        follow_up_required="Yes",
        follow_up_due_date="2026-09-08",
        linked_8d_record_id="SYN-8D-001",
        notes="Synthetic alarm.",
    )
    md = render_maintenance_markdown(record, redactor=Redactor())
    assert "Mitsubishi" not in md
    assert "drive_oem_a" in md


def test_ingest_maintenance_xlsx_writes_markdown(
    minimal_maintenance_xlsx: Path, tmp_path: Path
):
    out_dir = tmp_path / "maint_md"
    written = ingest_maintenance_xlsx(
        minimal_maintenance_xlsx, out_dir, redactor=Redactor()
    )
    assert len(written) == 1
    assert (out_dir / "ML-TEST-0001.md").exists()


# =============================================================================
# S3 — OEM PDF section extraction
# =============================================================================


def test_extract_pdf_sections_handles_owner_encrypted(tmp_path: Path):
    """The competition PDF is owner-encrypted; ingest_oem_pdf must decrypt with an
    empty password before extracting any troubleshooting section."""

    pypdf = pytest.importorskip("pypdf")
    if not _have_cryptography():
        pytest.skip("cryptography not installed")

    pdf_path = tmp_path / "manual.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.add_blank_page(width=300, height=400)
    page = writer.pages[0]
    page.merge_page(_text_page(pypdf, "1 Installation\n\nMount the unit per OEM drive manual."))
    page2 = writer.pages[1]
    page2.merge_page(
        _text_page(
            pypdf,
            "2 Troubleshooting\n\nAlarm 53 indicates drive-side vibration; reset the unit and inspect bearings.",
        )
    )
    writer.encrypt(user_password="", owner_password="", use_128bit=True)
    with pdf_path.open("wb") as fh:
        writer.write(fh)

    out_dir = tmp_path / "pdf_md"
    written = ingest_oem_pdf(pdf_path, out_dir, redactor=Redactor())
    assert len(written) >= 1
    joined = "\n".join(p.read_text(encoding="utf-8") for p in written)
    # Troubleshooting sections must surface after redaction.
    assert "drive_oem_a" in joined or "Troubleshooting" in joined
    # The brand tokens in the body must have been redacted.
    assert "Mitsubishi" not in joined


def _have_cryptography() -> bool:
    try:
        import cryptography  # noqa: F401
    except ImportError:
        return False
    return True


def _text_page(pypdf, text: str):
    """Return a pypdf Page object whose /Contents writes ``text`` as a Tj operand."""

    from pypdf.generic import (
        ArrayObject,
        DecodedStreamObject,
        DictionaryObject,
        FloatObject,
        NameObject,
    )

    content = DecodedStreamObject()
    content.set_data(
        b"BT /F1 12 Tf 50 750 Td (" + text.encode("latin-1", errors="replace") + b") Tj ET"
    )
    page = pypdf.PageObject(pypdf)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): DictionaryObject()})}
    )
    page[NameObject("/Contents")] = content
    page[NameObject("/MediaBox")] = ArrayObject(
        [FloatObject(0), FloatObject(0), FloatObject(612), FloatObject(792)]
    )
    return page


# =============================================================================
# S4 — index.json enrichment
# =============================================================================


def test_update_index_json_preserves_existing_entries_and_appends_drive_entries(
    minimal_index_json: Path,
    minimal_8d_xlsx: Path,
    minimal_maintenance_xlsx: Path,
):
    eight_d_records = parse_8d_workbook(minimal_8d_xlsx)
    maint_records = parse_maintenance_workbook(minimal_maintenance_xlsx)
    added = update_index_json(
        minimal_index_json,
        eight_d_records=eight_d_records,
        maintenance_records=maint_records,
    )
    assert added == 3  # 2 8D + 1 maintenance

    raw = json.loads(minimal_index_json.read_text(encoding="utf-8"))
    ids = [doc["document_id"] for doc in raw["documents"]]
    # Existing entry preserved.
    assert "manuals/cnc_operating_manual.md" in ids
    # Drive entries appended.
    assert "drive/past_8d_reports/8D_SYN-8D-001.md" in ids
    assert "drive/maint_logs/ML-TEST-0001.md" in ids

    drive_8d = next(
        d for d in raw["documents"] if d["document_id"] == "drive/past_8d_reports/8D_SYN-8D-001.md"
    )
    assert drive_8d["critical_citations"][0]["alarm_code"] == "53"
    assert drive_8d["critical_citations"][0]["symptom_key"] == "spindle_vibration"
    assert "tool_wear" not in drive_8d["key_topics"]  # 53 maps to vibration, not wear
    assert "spindle_vibration" in drive_8d["key_topics"]


def test_update_index_json_is_idempotent(
    minimal_index_json: Path,
    minimal_8d_xlsx: Path,
    minimal_maintenance_xlsx: Path,
):
    eight_d_records = parse_8d_workbook(minimal_8d_xlsx)
    maint_records = parse_maintenance_workbook(minimal_maintenance_xlsx)

    update_index_json(
        minimal_index_json,
        eight_d_records=eight_d_records,
        maintenance_records=maint_records,
    )
    first = json.loads(minimal_index_json.read_text(encoding="utf-8"))
    update_index_json(
        minimal_index_json,
        eight_d_records=eight_d_records,
        maintenance_records=maint_records,
    )
    second = json.loads(minimal_index_json.read_text(encoding="utf-8"))

    assert len(first["documents"]) == len(second["documents"])
    assert [d["document_id"] for d in first["documents"]] == [
        d["document_id"] for d in second["documents"]
    ]


# =============================================================================
# S5 — image ingest
# =============================================================================


def test__is_image_file_accepts_png_signature_without_extension(tmp_path: Path):
    png_path = tmp_path / "Error_part_99"
    png_path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 16)
    assert _is_image_file(png_path)


def test__is_image_file_accepts_jpg_extension(tmp_path: Path):
    jpg = tmp_path / "orig.jpg"
    jpg.write_bytes(b"\xff\xd8\xff\xe0\x00\x10JFIF")
    assert _is_image_file(jpg)


def test__is_image_file_rejects_text(tmp_path: Path):
    txt = tmp_path / "readme.txt"
    txt.write_text("hello")
    assert not _is_image_file(txt)


def test__classify_image_returns_normal_for_originals():
    category, severity = _classify_image(
        Path("Original Image/20261012_224035.jpg"), error_part_index=None
    )
    assert category == "normal"
    assert severity == "nominal"


def test__classify_image_returns_defect_for_error_parts():
    category, severity = _classify_image(
        Path("image error/Error_part_3"), error_part_index=3
    )
    assert severity in {"low", "medium", "high"}
    assert category in {"surface_scratch", "crack", "fracture", "assembly_gap"}


def test_ingest_competition_images_copies_and_appends(
    image_gdrive_layout: Path,
    tmp_path: Path,
):
    catalog_path = tmp_path / "catalog.json"
    catalog_path.write_text(json.dumps([{"id": "preset-im05-normal", "stays": True}]))

    destinations = (tmp_path / "web_pub", tmp_path / "mobile_pub")
    new_entries = ingest_competition_images(
        image_gdrive_layout,
        catalog_path=catalog_path,
        destinations=destinations,
    )

    # 2 originals + 2 defect images (one flat PNG + one nested jpg) = 4 new entries.
    assert len(new_entries) == 4

    # Files copied to both destinations.
    for dest in destinations:
        copied = list(dest.glob("competition_*"))
        assert len(copied) == 4

    # Catalog preserved the pre-existing preset and now has the new ones.
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    assert catalog[0]["id"] == "preset-im05-normal"
    competition_ids = [e["id"] for e in catalog if e["id"].startswith("preset-competition-")]
    assert len(competition_ids) == 4


def test_ingest_competition_images_is_idempotent(
    image_gdrive_layout: Path,
    tmp_path: Path,
):
    catalog_path = tmp_path / "catalog.json"
    catalog_path.write_text("[]")

    destinations = (tmp_path / "web_pub",)
    ingest_competition_images(
        image_gdrive_layout,
        catalog_path=catalog_path,
        destinations=destinations,
    )
    first = json.loads(catalog_path.read_text(encoding="utf-8"))
    ingest_competition_images(
        image_gdrive_layout,
        catalog_path=catalog_path,
        destinations=destinations,
    )
    second = json.loads(catalog_path.read_text(encoding="utf-8"))

    # Re-running must not double-insert; same total entry count.
    assert len(first) == len(second) == 4


# =============================================================================
# End-to-end pipeline (no fixtures -> minimal run_sync over a temp layout)
# =============================================================================


def test_run_sync_end_to_end_over_tmp_layout(tmp_path: Path, monkeypatch):
    """Smoke-test run_sync against a self-contained Drive cache under tmp_path."""

    from scripts import sync_gdrive_kb

    # Copy the real Drive fixtures into tmp_path so we don't touch the actual cache.
    real_drive = _REPO_ROOT / ".cache" / "gdrive" / "knowledge_base"
    if not (real_drive / "8D_Report_Historical_Checklist_for_AI_Development.xlsx").exists():
        pytest.skip("Real Drive cache missing; cannot run end-to-end smoke test.")

    tmp_drive = tmp_path / "drive_cache"
    shutil.copytree(real_drive, tmp_drive)

    kb_root = tmp_path / "kb"
    drive_output = kb_root / "drive"
    kb_root.mkdir(parents=True, exist_ok=True)
    # Copy the real curated index.json so S4 enrichment has something to extend.
    real_index = _REPO_ROOT / "data" / "knowledge_base" / "index.json"
    if real_index.exists():
        shutil.copy(real_index, kb_root / "index.json")
    catalog = tmp_path / "catalog.json"
    catalog.write_text("[]")
    destinations = (tmp_path / "web_pub", tmp_path / "mobile_pub")

    summary = sync_gdrive_kb.run_sync(
        gdrive_kb=tmp_drive,
        kb_root=kb_root,
        drive_output=drive_output,
        catalog_path=catalog,
        destinations=destinations,
    )

    outputs = summary["outputs"]
    assert outputs["8d_markdown"], "8D markdown list should be non-empty"
    assert outputs["maintenance_markdown"], "maintenance markdown list should be non-empty"
    # The OEM PDF is 554 pages; we expect several sections.
    assert len(outputs["pdf_markdown"]) >= 1
    # 10 originals + 7 Error_part files (flat PNGs) = 17 competition presets.
    assert outputs["competition_presets"] == 17
    # Redaction actually fired.
    assert summary["redaction_substitutions"].get("Mitsubishi", 0) > 0

    # Files copied to both web + mobile public dirs.
    for dest in destinations:
        copied = list(dest.glob("competition_*"))
        assert len(copied) == 17

    # index.json now contains at least one drive/* entry.
    index_raw = json.loads((kb_root / "index.json").read_text(encoding="utf-8"))
    drive_entries = [
        d for d in index_raw["documents"] if d["document_id"].startswith("drive/")
    ]
    assert drive_entries, "index.json should contain drive/* entries"