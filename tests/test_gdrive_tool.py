"""Unit tests for cloud-native GoogleDriveConnector and format-agnostic document parsing."""

from __future__ import annotations

import asyncio
import io
from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import openpyxl
import pytest

from backend.tools.gdrive_tool import (
    MIME_DOCX,
    MIME_GOOGLE_SHEET,
    MIME_XLSX,
    GoogleDriveConnector,
)


@pytest.fixture
def mock_connector(tmp_path: Path) -> GoogleDriveConnector:
    """Fixture providing a GoogleDriveConnector pointed to temporary cache."""
    return GoogleDriveConnector(
        folder_id="test-folder-123",
        cache_dir=tmp_path / "cache",
        timeout_seconds=5.0,
    )


def test_connector_initialization(mock_connector: GoogleDriveConnector, tmp_path: Path):
    """Verifies connector initialization and default attribute values."""
    assert mock_connector.folder_id == "test-folder-123"
    assert mock_connector.cache_dir == tmp_path / "cache"
    assert mock_connector.timeout == 5.0


def test_connector_availability_success(mock_connector: GoogleDriveConnector):
    """Verifies is_available returns True when Google Drive v3 client initializes."""
    with patch("google.auth.default", return_value=(MagicMock(), "test-project")):
        with patch("googleapiclient.discovery.build") as mock_build:
            mock_build.return_value = MagicMock()
            mock_connector._initialized = False
            assert mock_connector.is_available() is True


def test_connector_availability_fallback_on_error(mock_connector: GoogleDriveConnector):
    """Verifies is_available returns False when credentials fail."""
    with patch("google.auth.default", side_effect=Exception("No credentials")):
        mock_connector._initialized = False
        assert mock_connector.is_available() is False


def test_list_documents_api_success(mock_connector: GoogleDriveConnector):
    """Verifies parsing of Google Drive files.list response into GoogleDriveDocument models."""
    mock_service = MagicMock()
    mock_service.files().list().execute.return_value = {
        "files": [
            {
                "id": "file-1",
                "name": "cnc_manual.docx",
                "mimeType": MIME_DOCX,
                "size": "10240",
                "modifiedTime": "2026-03-12T10:00:00Z",
                "webViewLink": "https://drive.google.com/file-1",
            },
            {
                "id": "folder-1",
                "name": "manuals",
                "mimeType": "application/vnd.google-apps.folder",
                "size": "0",
                "modifiedTime": "2026-03-12T09:00:00Z",
            },
        ]
    }

    with patch.object(mock_connector, "_get_service", return_value=mock_service):
        docs = mock_connector.list_documents(folder_id="test-folder-123")

        assert len(docs) == 2
        doc1 = docs[0]
        assert doc1.file_id == "file-1"
        assert doc1.name == "cnc_manual.docx"
        assert doc1.size_bytes == 10240
        assert doc1.is_dir is False
        assert doc1.web_view_link == "https://drive.google.com/file-1"

        doc2 = docs[1]
        assert doc2.is_dir is True


def test_list_documents_fallback_when_service_none(mock_connector: GoogleDriveConnector):
    """Verifies list_documents gracefully falls back to mock documents when service is None."""
    with patch.object(mock_connector, "_get_service", return_value=None):
        docs = mock_connector.list_documents()
        assert len(docs) > 0
        assert any("cnc" in d.name.lower() for d in docs)


def test_stream_document_bytes(mock_connector: GoogleDriveConnector):
    """Verifies streaming raw bytes from Drive API get_media."""
    mock_service = MagicMock()
    mock_service.files().get_media().execute.return_value = b"raw file binary stream"

    with patch.object(mock_connector, "_get_service", return_value=mock_service):
        data = mock_connector.stream_document_bytes("file-123")
        assert data == b"raw file binary stream"


def test_export_google_doc_text(mock_connector: GoogleDriveConnector):
    """Verifies exporting native Google Docs and Sheets to plain text/CSV."""
    mock_service = MagicMock()
    mock_service.files().export_media().execute.return_value = b"Exported sheet row1,row2"

    with patch.object(mock_connector, "_get_service", return_value=mock_service):
        text = mock_connector.export_google_doc_text("sheet-123", MIME_GOOGLE_SHEET)
        assert "Exported sheet" in text


def test_parse_docx_in_memory():
    """Verifies in-memory extraction of headings, paragraphs, and tables from .docx bytes."""
    # Create an in-memory docx
    doc = docx.Document()
    doc.add_heading("CNC VMC-850 Operating Manual", level=1)
    doc.add_paragraph("Insert WNMG-080408 should be torqued to 4.2 N·m.")

    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Tool"
    table.rows[0].cells[1].text = "Insert"
    table.rows[0].cells[2].text = "Torque"
    table.rows[1].cells[0].text = "T07"
    table.rows[1].cells[1].text = "WNMG-080408"
    table.rows[1].cells[2].text = "4.2 N·m"

    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    connector = GoogleDriveConnector()
    extracted = connector.extract_document_text(
        file_bytes=docx_bytes, mime_type=MIME_DOCX, filename="manual.docx"
    )

    assert "CNC VMC-850 Operating Manual" in extracted
    assert "4.2 N·m" in extracted
    assert "| Tool | Insert | Torque |" in extracted
    assert "| T07 | WNMG-080408 | 4.2 N·m |" in extracted


def test_parse_xlsx_in_memory():
    """Verifies in-memory extraction of Excel sheets into header-anchored key-value lines."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tooling"
    ws.append(["Mã dao", "Mảnh cắt", "Lực siết", "Độ đảo"])
    ws.append(["T07", "WNMG-080408", "4.2 N·m", "0.003 mm"])
    ws.append(["T08", "CNMG-120408", "5.0 N·m", "0.004 mm"])

    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    connector = GoogleDriveConnector()
    extracted = connector.extract_document_text(
        file_bytes=xlsx_bytes, mime_type=MIME_XLSX, filename="maint.xlsx"
    )

    assert "### Sheet: Tooling" in extracted
    assert "Mã dao: T07" in extracted
    assert "Lực siết: 4.2 N·m" in extracted
    assert "Độ đảo: 0.003 mm" in extracted


def test_search_exact_documents(mock_connector: GoogleDriveConnector):
    """Verifies exact document search queries Drive API with fullText contains."""
    mock_service = MagicMock()
    mock_service.files().list().execute.return_value = {
        "files": [
            {
                "id": "file-search-1",
                "name": "cnc_tool_wear.docx",
                "mimeType": MIME_DOCX,
                "size": "5000",
                "modifiedTime": "2026-03-12T10:00:00Z",
                "webViewLink": "https://drive.google.com/file-search-1",
            }
        ]
    }

    with patch.object(mock_connector, "_get_service", return_value=mock_service):
        results = mock_connector.search_exact_documents(
            query_terms=["WNMG-080408", "CNC-02"], folder_id="test-folder"
        )

        assert len(results) == 1
        assert results[0].file_id == "file-search-1"
        assert results[0].name == "cnc_tool_wear.docx"


def test_fetch_to_cache(mock_connector: GoogleDriveConnector, tmp_path: Path):
    """Verifies fetching remote document into local cache directory without re-downloading."""
    content = b"Cached document content"

    with patch.object(mock_connector, "stream_document_bytes", return_value=content) as mock_stream:
        cached_path = mock_connector.fetch_to_cache("doc-1", "guide.md")
        assert cached_path.is_file()
        assert cached_path.read_bytes() == content
        mock_stream.assert_called_once()

        # Second call reuses cached file
        mock_stream.reset_mock()
        reused_path = mock_connector.fetch_to_cache("doc-1", "guide.md")
        assert reused_path == cached_path
        mock_stream.assert_not_called()


def test_async_connector_methods(mock_connector: GoogleDriveConnector):
    """Verifies async connector methods run asynchronously."""

    async def _run():
        with patch.object(mock_connector, "list_documents", return_value=[]):
            docs = await mock_connector.list_documents_async()
            assert docs == []

        with patch.object(mock_connector, "stream_document_bytes", return_value=b"async data"):
            data = await mock_connector.stream_document_bytes_async("doc-1")
            assert data == b"async data"

        with patch.object(mock_connector, "extract_document_text", return_value="async text"):
            text = await mock_connector.extract_document_text_async(b"async data")
            assert text == "async text"

    asyncio.run(_run())
